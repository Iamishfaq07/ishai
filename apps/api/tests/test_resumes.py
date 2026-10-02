"""Resume maker: the ATS score, CRUD and ownership, AI import and one-change-at-a-time suggestions."""

import json
from typing import Any

import pytest
from fastapi.testclient import TestClient

from orbit_api.providers.llm.base import LLMResponse, Usage
from orbit_api.resumes import ai
from orbit_api.resumes.ats import AtsCheck, date_key, date_style, extract_keywords, match_keywords, score_resume
from orbit_api.resumes.schema import (
    Basics,
    Education,
    Experience,
    ExtraSection,
    Link,
    ResumeData,
    ResumeDesign,
    SkillGroup,
)
from tests.test_api import _register

JOB = """Senior Backend Engineer. We're looking for an engineer with strong Python and PostgreSQL experience to build
scalable microservices on AWS. You will design REST and GraphQL APIs, run services on Kubernetes and Docker, and own
CI/CD with GitHub Actions. Experience with Kafka, Redis and observability (Datadog) is a plus."""


def strong_resume() -> ResumeData:
    return ResumeData(
        basics=Basics(name="Sam Lee", email="sam@example.com", phone="+1 555 123 4567", location="Austin, TX",
                      links=[Link(label="LinkedIn", url="linkedin.com/in/samlee")],
                      summary="Backend engineer with 7 years building payment systems in Python and PostgreSQL on "
                              "AWS. Designs scalable microservices and REST APIs and leads teams through large "
                              "migrations. Looking for a senior backend role on a platform team."),
        experience=[
            Experience(title="Senior Engineer", company="PayCo", start="Jan 2020", end="", current=True, bullets=[
                "Led the migration of the ledger to event sourcing on Kafka, cutting reconciliation time by 60%",
                "Built 12 Python microservices on Kubernetes and Docker serving 3M transactions a day",
                "Designed GraphQL and REST APIs used by 40 internal teams with 99.99% uptime",
                "Reduced p95 latency from 800 ms to 120 ms with Redis caching and PostgreSQL query tuning"]),
            Experience(title="Engineer", company="Shop Inc", start="Jun 2016", end="Dec 2019", bullets=[
                "Shipped CI/CD with GitHub Actions, cutting release time from 2 days to 30 minutes",
                "Mentored 4 junior engineers and ran the team's system design reviews every quarter"]),
        ],
        education=[Education(school="UT Austin", degree="BSc", field="Computer Science", start="2012", end="2016")],
        skills=[SkillGroup(name="Languages", items=["Python", "SQL", "Go"]),
                SkillGroup(name="Infrastructure", items=["AWS", "Kubernetes", "Docker", "Terraform", "Kafka",
                                                         "Redis", "PostgreSQL", "Datadog"])],
    )


# ---- ATS score ---------------------------------------------------------------------------------------------


def test_weak_resume_scores_low_with_a_fix_for_every_missed_point() -> None:
    weak = ResumeData(basics=Basics(name="Sam"),
                      experience=[Experience(title="Engineer", company="Acme", bullets=["I worked on the backend"])])
    report = score_resume(weak, job_text=JOB)
    assert report.score < 30 and report.verdict == "Weak" and report.has_job
    missed = [c for c in report.checks if not c.passed]
    assert missed and all(c.fix for c in missed if c.id != "keywords" or report.keywords)
    by_id = {c.id: c for c in report.checks}
    assert by_id["email"].target == "basics.email"
    assert by_id["action_verbs"].target == f"experience.{weak.experience[0].id}.bullets.0"
    assert by_id["no_pronouns"].ai_fixable and not by_id["no_pronouns"].passed
    assert sum(c.max_points for c in report.checks) == 100


def test_strong_resume_scores_high_and_matches_the_job() -> None:
    report = score_resume(strong_resume(), job_text=JOB)
    assert report.score >= 85 and report.verdict == "Strong"
    assert report.keywords is not None and "Datadog" in report.keywords.matched
    assert not report.keywords.missing or report.keywords.missing == ["Observability"]


def test_without_a_job_post_keywords_are_left_out_and_the_rest_scaled_to_100() -> None:
    report = score_resume(strong_resume())
    assert not report.has_job and report.keywords is None
    assert "keywords" not in {c.category for c in report.checks}
    assert sum(c.max_points for c in report.checks) == 75 and report.score >= 85


def test_layout_photo_and_two_columns_cost_points() -> None:
    plain = score_resume(strong_resume()).score
    styled = score_resume(strong_resume(), ResumeDesign(template="sidebar", show_photo=True,
                                                        photo="data:image/png;base64,iVBORw0KGgo="))
    assert styled.score < plain
    assert {c.id for c in styled.checks if not c.passed} >= {"single_column", "no_photo"}


def test_dates_and_keywords_helpers() -> None:
    assert [date_style(d) for d in ("Jan 2021", "01/2021", "2021", "Present", "", "spring")] == [
        "month_year", "numeric", "year", None, None, "other"]
    mixed = strong_resume()
    mixed.experience[1].start = "06/2016"
    assert not {c.id: c for c in score_resume(mixed).checks}["dates"].passed
    kws = extract_keywords(JOB)
    assert {"Python", "Kubernetes", "CI/CD", "Datadog"} <= set(kws)  # written the way the post writes them
    assert not {"CI", "GitHub", "Backend", "APIs"} & set(kws)  # parts of longer terms, titles, generic words
    got = match_keywords(["node.js", "a/b testing", "c++", "microservice"], ResumeData(
        basics=Basics(summary="Nodejs services, A/B testing, C++ and microservices")))
    assert got.missing == [] and len(got.matched) == 4


def test_design_is_validated() -> None:
    with pytest.raises(ValueError):
        ResumeDesign(accent="red")
    with pytest.raises(ValueError):
        ResumeDesign(photo="javascript:alert(1)")
    assert ResumeDesign(accent="#ABCDEF").accent == "#abcdef"
    assert ResumeData(section_order=["skills"]).section_order[0] == "skills"  # the rest keep their places


def _checks(data: ResumeData) -> dict[str, AtsCheck]:
    return {c.id: c for c in score_resume(data).checks}


def test_unfilled_placeholders_are_caught_before_they_print() -> None:
    assert _checks(strong_resume())["placeholders"].passed
    data = strong_resume()
    data.experience[1].bullets[0] = "Shipped CI/CD with GitHub Actions, cutting release time by [X%]"
    check = _checks(data)["placeholders"]
    assert not check.passed and "[X%]" in (check.fix or "") and check.ai_fixable  # the AI rewrites it finished
    assert check.target == f"experience.{data.experience[1].id}.bullets.0"


def test_repeated_opening_verbs_cost_points_and_suggest_others() -> None:
    assert _checks(strong_resume())["varied_verbs"].passed
    data = strong_resume()
    role = data.experience[0]
    role.bullets = [f"Led the {x} migration for the payments ledger, cutting costs by 20%" for x in "ABCD"]
    check = _checks(data)["varied_verbs"]
    assert not check.passed and check.ai_fixable
    assert "4 bullets start with \"Led\"" in (check.fix or "") and "Directed" in (check.fix or "")
    assert check.target == f"experience.{role.id}.bullets.2"  # the first one past the two allowed


def test_cliches_are_flagged_with_where_they_are() -> None:
    assert _checks(strong_resume())["cliches"].passed
    data = strong_resume()
    data.basics.summary += " A results-driven team player."
    check = _checks(data)["cliches"]
    assert not check.passed and check.target == "basics.summary" and check.ai_fixable
    assert "results-driven" in (check.fix or "")


def test_roles_must_be_most_recent_first_with_dates_in_order() -> None:
    assert _checks(strong_resume())["date_order"].passed
    backwards = strong_resume()
    backwards.experience.reverse()
    check = _checks(backwards)["date_order"]
    assert not check.passed and check.target == f"experience.{backwards.experience[1].id}"
    inverted = strong_resume()
    inverted.experience[1].start, inverted.experience[1].end = "Dec 2019", "Jun 2016"
    check = _checks(inverted)["date_order"]
    assert not check.passed and "ends before it starts" in (check.fix or "")
    overlapping = strong_resume()  # a longer role that started earlier, listed second: not out of order
    overlapping.experience[1].start, overlapping.experience[1].end = "Jun 2016", "Dec 2024"
    assert _checks(overlapping)["date_order"].passed
    assert [date_key(v) for v in ("Jan 2021", "03/2020", "2019-07", "2018", "Present", "soon")] == [
        (2021, 1), (2020, 3), (2019, 7), (2018, 1), None, None]
    assert date_key("2018", end=True) == (2018, 12) and date_key("Present", end=True) == (9999, 12)


def test_extra_sections_keep_their_place_and_are_read_and_exported(client: TestClient) -> None:
    langs = ExtraSection(id="lang1", title="Languages", items=["English (native)", "Spanish (fluent)"],
                         style="inline")
    awards = ExtraSection(id="aw1", title="Awards", items=["Hackathon winner, PyCon 2022"])
    data = strong_resume()
    data.extras = [langs, awards]
    data.section_order = ["experience", "extra:aw1", "skills", "extra:gone"]  # unknown entries are dropped
    data = ResumeData.model_validate(data.model_dump())
    assert data.section_order[:3] == ["experience", "extra:aw1", "skills"] and "extra:lang1" in data.section_order
    assert "extra:gone" not in data.section_order and len(data.section_order) == 8
    with pytest.raises(ValueError):
        ResumeData(extras=[ExtraSection(title=str(i)) for i in range(6)])  # at most 5

    from orbit_api.documents.parsing import extract_text

    h = _register(client)
    r = client.post("/v1/resumes/docx", headers=h, json={"data": data.model_dump(mode="json")})
    lines = extract_text(r.content, "docx").split("\n")
    assert lines.index("EXPERIENCE") < lines.index("AWARDS") < lines.index("SKILLS")
    assert "Hackathon winner, PyCon 2022" in lines and "English (native)  ·  Spanish (fluent)" in lines
    assert "Spanish" in match_keywords(["Spanish"], data).matched


def test_import_reads_extra_sections() -> None:
    raw = {"basics": {"name": "Ana"}, "extras": [
        {"title": "Languages", "items": ["English", "Portuguese"]}, {"title": "Volunteering", "items": ["Mentor"]},
        {"title": "", "items": ["no title"]}, {"title": "Empty", "items": []}, "junk"]}
    extras = ai.parse_structured(json.dumps(raw)).extras
    assert [(x.title, x.style) for x in extras] == [("Languages", "inline"), ("Volunteering", "bullets")]


# ---- CRUD & ownership --------------------------------------------------------------------------------------


def test_crud_is_scoped_to_the_owner(client: TestClient) -> None:
    h = _register(client)
    other = _register(client)
    body = {"title": "Backend", "data": strong_resume().model_dump(mode="json"), "job_text": JOB,
            "design": {"template": "modern", "accent": "#0f766e", "font": "georgia"}}
    created = client.post("/v1/resumes", headers=h, json=body)
    assert created.status_code == 201, created.text
    rid = created.json()["id"]
    assert created.json()["design"]["template"] == "modern"

    listed = client.get("/v1/resumes", headers=h).json()
    assert listed[0]["id"] == rid and listed[0]["score"] >= 85 and listed[0]["has_job"]
    assert listed[0]["name"] == "Sam Lee"

    for method in ("GET", "PUT", "DELETE"):
        r = client.request(method, f"/v1/resumes/{rid}", headers=other, json=body if method == "PUT" else None)
        assert r.status_code == 404, method
    assert client.get("/v1/resumes", headers=other).json() == []

    body["title"] = "Backend v2"
    assert client.put(f"/v1/resumes/{rid}", headers=h, json=body).json()["title"] == "Backend v2"
    assert client.post("/v1/resumes", headers=h, json={**body, "design": {"accent": "blue"}}).status_code == 422
    assert client.delete(f"/v1/resumes/{rid}", headers=h).status_code == 204
    assert client.get(f"/v1/resumes/{rid}", headers=h).status_code == 404


def test_score_endpoint_scores_a_draft(client: TestClient) -> None:
    h = _register(client)
    r = client.post("/v1/resumes/score", headers=h, json={"data": strong_resume().model_dump(mode="json"),
                                                          "job_text": JOB})
    assert r.status_code == 200 and r.json()["score"] >= 85 and r.json()["keywords"]["matched"]
    assert client.post("/v1/resumes/score", json={"data": {}}).status_code == 401


def test_export_includes_resumes(client: TestClient) -> None:
    h = _register(client)
    client.post("/v1/resumes", headers=h, json={"title": "Mine"})
    assert [r["title"] for r in client.get("/v1/auth/me/export", headers=h).json()["resumes"]] == ["Mine"]


# ---- Word export -------------------------------------------------------------------------------------------


def test_word_export_reads_back_in_order(client: TestClient) -> None:
    import docx

    from orbit_api.documents.parsing import extract_text

    h = _register(client)
    data = strong_resume()
    data.section_order = ["experience", "skills", "education", "summary", "projects", "certifications"]
    r = client.post("/v1/resumes/docx", headers=h, json={
        "data": data.model_dump(mode="json"), "design": {"font": "georgia", "accent": "#0f766e", "page": "letter"}})
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.wordprocessingml")
    text = extract_text(r.content, "docx")
    lines = text.split("\n")
    assert lines[0] == "Sam Lee" and "sam@example.com" in lines[1]
    # The chosen section order, empty sections left out, the job line with its dates after a tab.
    heads = [line for line in lines if line in ("EXPERIENCE", "SKILLS", "EDUCATION", "SUMMARY", "PROJECTS")]
    assert heads == ["EXPERIENCE", "SKILLS", "EDUCATION", "SUMMARY"]
    assert "Senior Engineer, PayCo\tJan 2020 – Present" in lines
    assert "Infrastructure: AWS, Kubernetes, Docker, Terraform, Kafka, Redis, PostgreSQL, Datadog" in lines
    d = docx.Document(__import__("io").BytesIO(r.content))
    assert d.styles["Normal"].font.name == "Georgia"
    assert round(d.sections[0].page_width.mm) == 216
    assert sum(p.style.name == "List Bullet" for p in d.paragraphs) == 6
    assert sum(p.style.name == "Heading 1" for p in d.paragraphs) == 4
    assert client.post("/v1/resumes/docx", json={"data": {}}).status_code == 401


# ---- AI import ---------------------------------------------------------------------------------------------

PASTED = """Priya Raman
priya@example.com | +44 20 7946 0958
- Led a team of 5 building the checkout redesign
- Cut page load time by 35% across the storefront
- Ran weekly experiments with the growth team"""


def test_import_from_pasted_text_file_and_document(client: TestClient) -> None:
    h = _register(client)
    r = client.post("/v1/resumes/import", headers=h, json={"source": "text", "text": PASTED})
    assert r.status_code == 200, r.text
    data = r.json()["data"]
    assert data["basics"]["name"] == "Priya Raman" and data["basics"]["email"] == "priya@example.com"
    assert r.json()["title"] == "Priya Raman's resume"
    assert data["experience"][0]["bullets"][0].startswith("Led a team")

    too_short = client.post("/v1/resumes/import", headers=h, json={"source": "linkedin", "text": "hi"})
    assert too_short.status_code == 422

    f = client.post("/v1/resumes/import-file", headers=h, files={"file": ("cv.txt", PASTED.encode(), "text/plain")})
    assert f.status_code == 200 and f.json()["data"]["basics"]["name"] == "Priya Raman"
    bad = client.post("/v1/resumes/import-file", headers=h, files={"file": ("cv.exe", b"MZ\x90\x00", "x/y")})
    assert bad.status_code == 422

    doc = client.post("/v1/documents/text", headers=h, json={"kind": "resume", "title": "CV", "text": PASTED}).json()
    from_doc = client.post("/v1/resumes/import", headers=h, json={"source": "document", "document_id": doc["id"]})
    assert from_doc.status_code == 200
    other = _register(client)
    assert client.post("/v1/resumes/import", headers=other,
                       json={"source": "document", "document_id": doc["id"]}).status_code == 404


def test_parse_structured_clips_and_drops_malformed_items() -> None:
    raw = {"basics": {"name": "A" * 500, "email": "a@b.co", "links": [{"url": "x.com"}, "junk", {"url": ""}]},
           "experience": [{"title": "Dev", "company": "Co", "bullets": ["one", "", 3, "x" * 900]}, "junk", {}],
           "skills": [{"name": "Tools", "items": ["Git"] * 60}, {"name": "Empty", "items": []}],
           "education": "nope"}
    data = ai.parse_structured("Here you go: " + json.dumps(raw))
    assert len(data.basics.name) == 120 and [link.url for link in data.basics.links] == ["x.com"]
    assert len(data.experience) == 1 and data.experience[0].bullets[:3] == ["one", "3", "x" * 400]
    assert len(data.skills) == 1 and len(data.skills[0].items) == 40 and data.education == []
    with pytest.raises(ValueError):
        ai.parse_structured(json.dumps({"basics": {}}))


# ---- AI suggestions ----------------------------------------------------------------------------------------


def test_suggestions_name_the_exact_text_they_replace(client: TestClient) -> None:
    h = _register(client)
    data = strong_resume().model_dump(mode="json")
    r = client.post("/v1/resumes/suggest", headers=h, json={"data": data, "mode": "improve"})
    assert r.status_code == 200, r.text
    fields = ai.editable_fields(strong_resume().model_validate(data))
    rewrites = [s for s in r.json()["suggestions"] if s["kind"] == "rewrite"]
    assert rewrites
    for s in rewrites:
        assert s["path"] in fields and s["before"] == fields[s["path"]][1] and s["after"] != s["before"]

    first_bullet = f"experience.{data['experience'][0]['id']}.bullets.0"
    focused = client.post("/v1/resumes/suggest", headers=h, json={"data": data, "focus": first_bullet}).json()
    assert [s["path"] for s in focused["suggestions"]] == [first_bullet]

    assert client.post("/v1/resumes/suggest", headers=h, json={"data": data, "mode": "tailor"}).status_code == 422
    assert client.post("/v1/resumes/suggest", headers=h,
                       json={"data": data, "focus": "basics.email"}).status_code == 422
    assert client.post("/v1/resumes/suggest", headers=h, json={"data": {}}).status_code == 422  # nothing to rewrite


def test_parse_suggestions_rejects_unknown_paths_no_ops_and_duplicates() -> None:
    data = strong_resume()
    path = f"experience.{data.experience[0].id}.bullets.0"
    same = data.experience[0].bullets[0]

    def resp(items: list[dict[str, Any]]) -> LLMResponse:
        return LLMResponse(text=json.dumps({"suggestions": items}), model="m", usage=Usage())

    got = ai.parse_suggestions(resp([
        {"path": "experience.nope.bullets.0", "after": "x"},
        {"path": path, "after": same},  # no change
        {"path": path, "after": "• Led the ledger migration to event sourcing, cutting reconciliation time by 60%",
         "reason": "Leads with the result."},
        {"path": path, "after": "duplicate"},
        {"path": "basics.summary", "after": "S" * 3000},
    ]), data)
    assert [s.path for s in got] == [path, "basics.summary"]
    assert got[0].after.startswith("Led ") and got[0].before == same  # bullet character stripped
    assert len(got[1].after) == 1_500


# ---- cover letter ------------------------------------------------------------------------------------------


def test_cover_letter_is_written_saved_and_kept_by_later_saves(client: TestClient) -> None:
    h = _register(client)
    data = strong_resume().model_dump(mode="json")
    r = client.post("/v1/resumes/cover-letter", headers=h, json={"data": data, "job_text": JOB})
    assert r.status_code == 200, r.text
    letter = r.json()["text"]
    assert letter.startswith("Dear Hiring Manager,") and letter.rstrip().endswith("Sam Lee")

    created = client.post("/v1/resumes", headers=h, json={"title": "Mine", "data": data, "cover_letter": letter})
    rid = created.json()["id"]
    assert created.json()["cover_letter"] == letter
    # A save that doesn't mention the letter (an older client) leaves it alone; an explicit one replaces it.
    kept = client.put(f"/v1/resumes/{rid}", headers=h, json={"title": "Mine", "data": data})
    assert kept.json()["cover_letter"] == letter
    edited = client.put(f"/v1/resumes/{rid}", headers=h, json={"title": "Mine", "data": data, "cover_letter": "Edited"})
    assert edited.json()["cover_letter"] == "Edited"

    empty = client.post("/v1/resumes/cover-letter", headers=h, json={"data": {}})
    assert empty.status_code == 422 and "experience" in empty.json()["detail"]


def test_cover_letter_word_file_has_the_header_and_paragraphs(client: TestClient) -> None:
    from orbit_api.documents.parsing import extract_text

    h = _register(client)
    text = "Dear Hiring Manager,\n\nFirst paragraph.\n\nSecond paragraph.\n\nSincerely,\nSam Lee"
    r = client.post("/v1/resumes/cover-letter/docx", headers=h,
                    json={"text": text, "data": strong_resume().model_dump(mode="json")})
    assert r.status_code == 200
    lines = extract_text(r.content, "docx").split("\n")
    assert lines[0] == "Sam Lee" and "sam@example.com" in lines[1]
    assert lines[3:] == ["Dear Hiring Manager,", "First paragraph.", "Second paragraph.", "Sincerely,", "Sam Lee"]


def test_parse_cover_letter_strips_wrappers_and_rejects_scraps() -> None:
    body = "Dear Hiring Manager,\n\n" + "A real paragraph about the work. " * 10 + "\n\n\n\nSincerely,\nAna"
    assert ai.parse_cover_letter("Here's your letter:\n```\n" + body + "\n```").startswith("Dear Hiring Manager,")
    assert "\n\n\n" not in ai.parse_cover_letter(body)
    with pytest.raises(ValueError):
        ai.parse_cover_letter("Sorry, I can't.")


def test_job_keywords_skip_fragments_perks_and_lookalike_words() -> None:
    from orbit_api.resumes.ats import extract_keywords

    job = ("Northwind builds payment rails for merchants. Remote or hybrid.\n"
           "Requirements\n- Experience with gRPC and PostgreSQL\n- Benefits include equity, health insurance and PTO\n")
    got = {k.lower() for k in extract_keywords(job)}
    assert {"grpc", "postgresql"} <= got
    assert not got & {"rpc", "rails", "remote", "hybrid", "pto"}  # "gRPC" isn't "g"+"RPC"; "rails" isn't Ruby on Rails
    assert "ruby on rails" in {k.lower() for k in extract_keywords("We use Ruby on Rails and PostgreSQL daily.")}
