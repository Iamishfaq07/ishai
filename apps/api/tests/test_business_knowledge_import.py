"""Policy extraction previews, scoped provenance, manual publication and read-only AI review."""

import io
import uuid
import zipfile
from collections.abc import Iterator
from concurrent.futures import ThreadPoolExecutor
from datetime import UTC, datetime, timedelta
from threading import Barrier
from typing import Any

import jwt
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import func, select

from orbit_api.auth.security import decode_access_token
from orbit_api.business.knowledge_import import PURPOSE, signing_key
from orbit_api.business.knowledge_import_parser import MAX_BYTES, clean_text, parse
from orbit_api.documents.parsing import ParseError
from orbit_api.models import BizKnowledgeDocument, BizKnowledgeEvent, BizMember, ProviderRequest
from orbit_api.providers.llm.mock import MockLLMProvider
from tests.business_helpers import auth, db_call, docx_bytes, make_client, new_company, uid
from tests.test_business_knowledge import act, content, member, path, publish


def _invalid_child_reply(data: bytes, fmt: str, sender: Any) -> None:
    sender.send_bytes(b"[]")
    sender.close()


def _stuck_child(data: bytes, fmt: str, sender: Any) -> None:
    import time

    time.sleep(5)


@pytest.fixture
def client(monkeypatch: pytest.MonkeyPatch) -> Iterator[TestClient]:
    yield from make_client(monkeypatch)


def imports(cid: str) -> str:
    return f"/v1/business/companies/{cid}/knowledge-imports"


def preview(
    client: TestClient,
    who: dict[str, str],
    cid: str,
    data: bytes | None = None,
    filename: str = "Interview policy.txt",
    mime: str = "text/plain",
) -> Any:
    return client.post(
        imports(cid) + "/preview",
        headers=auth(who),
        files={"file": (filename, data if data is not None else content()["body"].encode(), mime)},
    )


def commit(
    client: TestClient,
    who: dict[str, str],
    cid: str,
    result: dict[str, Any],
    request_id: str | None = None,
    **extra: Any,
) -> Any:
    return client.post(
        imports(cid) + "/commit",
        headers=auth(who),
        json={
            "request_id": request_id or str(uuid.uuid4()),
            "receipt": result["receipt"],
            "data": content(body=result["text"]),
            "reviewed_original": True,
            **extra,
        },
    )


def review(client: TestClient, who: dict[str, str], cid: str, result: dict[str, Any]) -> Any:
    return client.post(
        imports(cid) + "/review",
        headers=auth(who),
        json={"receipt": result["receipt"], "data": content(body=result["text"])},
    )


def test_preview_saves_nothing_and_commit_is_unpublished_with_private_provenance(client: TestClient) -> None:
    owner, company, _ = new_company(client)
    cid = company["id"]
    reader = member(client, cid, "interviewer")
    response = preview(client, owner, cid)
    assert response.status_code == 200, response.text
    result = response.json()
    assert result["source"]["format"] == "txt" and result["text"] == content()["body"]
    assert client.get(path(cid), headers=auth(owner)).json()["items"] == []
    saved = commit(client, owner, cid, result)
    assert saved.status_code == 200, saved.text
    row = saved.json()
    assert row["published"] is None and row["published_revision"] is None
    assert row["import_source"]["filename"] == "Interview policy.txt"
    assert not row["import_source"]["edited_before_save"]
    duplicate = preview(client, owner, cid).json()
    assert duplicate["duplicates"] == [{"id": row["id"], "title": row["data"]["title"], "published": False}]
    assert client.get(path(cid), headers=auth(reader)).json()["items"] == []
    row = publish(client, owner, cid, row)
    read = client.get(path(cid) + "/" + row["id"], headers=auth(reader))
    assert read.status_code == 200 and read.json()["import_source"] is None
    assert "Interview policy.txt" not in read.text and "file_sha256" not in read.text
    assert act(client, owner, cid, row, "delete").status_code == 200

    async def events(db: Any) -> int:
        return int(
            await db.scalar(
                select(func.count()).select_from(BizKnowledgeEvent).where(BizKnowledgeEvent.document_id == row["id"])
            )
            or 0
        )

    assert db_call(client, events) == 0


def test_receipt_is_scoped_expiring_not_a_session_and_cannot_be_tampered_with(client: TestClient) -> None:
    owner, company, _ = new_company(client)
    cid = company["id"]
    admin = member(client, cid, "admin")
    other, other_company, _ = new_company(client)
    result = preview(client, owner, cid).json()
    assert decode_access_token(result["receipt"]) is None
    assert commit(client, admin, cid, result).status_code == 409
    assert commit(client, other, other_company["id"], result).status_code == 409
    assert commit(client, owner, cid, {**result, "receipt": result["receipt"][:-8] + "tampered"}).status_code == 409
    claims = jwt.decode(result["receipt"], signing_key(), algorithms=["HS256"], audience=PURPOSE)
    claims["iat"] = datetime.now(UTC) - timedelta(hours=2)
    claims["exp"] = datetime.now(UTC) - timedelta(hours=1)
    expired = jwt.encode(claims, signing_key(), algorithm="HS256")
    assert commit(client, owner, cid, {**result, "receipt": expired}).status_code == 409
    assert commit(client, owner, cid, result, reviewed_original=False).status_code == 422
    assert commit(client, owner, cid, result, reviewed_original="true").status_code == 422
    assert commit(client, owner, cid, result, published=True).status_code == 422


def test_owner_permissions_are_rechecked_after_extraction(client: TestClient, monkeypatch: pytest.MonkeyPatch) -> None:
    from orbit_api.business import knowledge_import
    from orbit_api.db import sessionmaker

    owner, company, _ = new_company(client)
    cid = company["id"]
    recruiter = member(client, cid, "recruiter")
    assert preview(client, recruiter, cid).status_code == 403
    original = knowledge_import.parse

    def demote(data: bytes, fmt: str) -> dict[str, Any]:
        result = original(data, fmt)

        async def change() -> None:
            async with sessionmaker()() as db:
                m = await db.scalar(
                    select(BizMember).where(BizMember.company_id == cid, BizMember.user_id == uid(owner))
                )
                m.role = "recruiter"
                await db.commit()

        client.portal.call(change)
        return result

    monkeypatch.setattr(knowledge_import, "parse", demote)
    assert preview(client, owner, cid).status_code == 403


def test_concurrent_commit_is_idempotent_and_records_reviewed_edits(client: TestClient) -> None:
    owner, company, _ = new_company(client)
    cid = company["id"]
    result = preview(client, owner, cid).json()
    request_id = str(uuid.uuid4())
    data = content(body=result["text"] + " Reviewed wording was clarified before saving.")
    barrier = Barrier(2)

    def send() -> Any:
        barrier.wait(timeout=10)
        return commit(client, owner, cid, result, request_id, data=data)

    with ThreadPoolExecutor(max_workers=2) as pool:
        responses = list(pool.map(lambda _: send(), range(2)))
    assert all(r.status_code == 200 for r in responses), [r.text for r in responses]
    assert responses[0].json()["id"] == responses[1].json()["id"] == request_id
    assert responses[0].json()["import_source"]["edited_before_save"]
    assert responses[0].json()["import_source"]["saved_text_sha256"] != result["source"]["extracted_sha256"]
    assert commit(client, owner, cid, result, request_id).status_code == 409

    async def count(db: Any) -> int:
        return int(
            await db.scalar(
                select(func.count()).select_from(BizKnowledgeDocument).where(BizKnowledgeDocument.company_id == cid)
            )
            or 0
        )

    assert db_call(client, count) == 1


@pytest.mark.parametrize(
    "data,filename,mime,expected",
    [
        (b"", "empty.txt", "text/plain", 413),
        (b"x" * (MAX_BYTES + 1), "large.txt", "text/plain", 413),
        (b"small text", "tiny.txt", "text/plain", 422),
        (b"PKnot a docx", "bad.docx", "application/octet-stream", 422),
        (b"not a pdf", "renamed.pdf", "application/pdf", 422),
        (b"x" * 20001, "large-text.md", "text/markdown", 422),
        (b"policy" * 20, "script.py", "text/plain", 422),
        (b"\xff\xfe" * 50, "wrong-encoding.txt", "text/plain", 422),
    ],
)
def test_unsupported_malformed_or_oversized_uploads_are_rejected(
    client: TestClient, data: bytes, filename: str, mime: str, expected: int
) -> None:
    owner, company, _ = new_company(client)
    result = preview(client, owner, company["id"], data, filename, mime)
    assert result.status_code == expected, result.text


def test_docx_worker_preserves_body_table_order_and_does_not_fetch_links() -> None:
    from docx import Document

    doc = Document()
    doc.add_paragraph("Interview policy: record independent observations before reading peer feedback.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Step", "Owner"
    table.cell(1, 0).text, table.cell(1, 1).text = "Submit scorecard", "Interviewer"
    doc.add_paragraph("Then meet the coordinator to discuss disagreements.")
    doc.sections[0].header.paragraphs[0].text = "PRIVATE HEADER NOT EXTRACTED"
    buf = io.BytesIO()
    doc.save(buf)
    result = parse(buf.getvalue(), "docx")
    assert result["text"].index("Interview policy") < result["text"].index("Step | Owner")
    assert result["text"].index("Submit scorecard | Interviewer") < result["text"].index("Then meet")
    assert "PRIVATE HEADER" not in result["text"]
    assert result["pages"] is None


def test_docx_archive_bomb_guard_and_text_normalization() -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "types")
        archive.writestr("word/document.xml", "x" * (2 * 1024 * 1024))
    with pytest.raises(ParseError, match="compression bomb"):
        parse(buffer.getvalue(), "docx")
    assert clean_text("\ufeffLine one\r\nLine\u202e two\x00") == "Line one\nLine two"


def test_malformed_worker_reply_and_timeout_fail_closed_and_release_capacity(monkeypatch: pytest.MonkeyPatch) -> None:
    from orbit_api.business import knowledge_import_parser

    data = docx_bytes(content()["body"])
    monkeypatch.setattr(knowledge_import_parser, "_worker", _invalid_child_reply)
    with pytest.raises(ParseError, match="invalid result"):
        parse(data, "docx")
    monkeypatch.setattr(knowledge_import_parser, "_worker", _stuck_child)
    monkeypatch.setattr(knowledge_import_parser, "TIMEOUT_S", 0.05)
    with pytest.raises(ParseError, match="timed out"):
        parse(data, "docx")
    monkeypatch.undo()
    assert parse(data, "docx")["text"] == content()["body"]


def test_pdf_worker_rejects_image_only_encrypted_and_excess_pages() -> None:
    from pypdf import PdfWriter

    for encrypted, count in ((False, 1), (True, 1), (False, 41)):
        pdf = PdfWriter()
        for _ in range(count):
            pdf.add_blank_page(width=400, height=400)
        if encrypted:
            pdf.encrypt("password")
        buf = io.BytesIO()
        pdf.write(buf)
        with pytest.raises(ParseError):
            parse(buf.getvalue(), "pdf")


def test_text_pdf_extracts_in_worker_with_page_count() -> None:
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    pdf = PdfWriter()
    page = pdf.add_blank_page(width=600, height=800)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): pdf._add_object(font)})}
    )
    stream = DecodedStreamObject()
    stream.set_data(
        b"BT /F1 12 Tf 50 700 Td (Interview policy: record independent scorecards before reading peer feedback.) Tj ET"
    )
    page[NameObject("/Contents")] = pdf._add_object(stream)
    buffer = io.BytesIO()
    pdf.write(buffer)
    result = parse(buffer.getvalue(), "pdf")
    assert result["pages"] == 1 and "record independent scorecards" in result["text"]


def test_ai_review_uses_quoted_excerpts_and_never_saves_a_document(client: TestClient) -> None:
    owner, company, _ = new_company(client)
    cid = company["id"]
    result = preview(client, owner, cid, (content()["body"] * 80).encode()).json()
    answer = review(client, owner, cid, result)
    assert answer.status_code == 200, answer.text
    data = answer.json()
    assert data["source"] == "mock" and data["excerpts_truncated"]
    assert len(data["facts"]) == 6
    facts = {f["id"]: f["text"] for f in data["facts"]}
    assert all(t["source_quote"] in facts[t["fact_id"]] for t in data["tips"])
    assert client.get(path(cid), headers=auth(owner)).json()["items"] == []


def test_ai_rechecks_permissions_after_generation(client: TestClient, monkeypatch: pytest.MonkeyPatch) -> None:
    from orbit_api.db import sessionmaker

    owner, company, _ = new_company(client)
    cid = company["id"]
    result = preview(client, owner, cid).json()
    original = MockLLMProvider.generate

    async def demote(self: Any, request: Any) -> Any:
        async with sessionmaker()() as db:
            m = await db.scalar(select(BizMember).where(BizMember.company_id == cid, BizMember.user_id == uid(owner)))
            m.role = "interviewer"
            await db.commit()
        return await original(self, request)

    monkeypatch.setattr(MockLLMProvider, "generate", demote)
    assert review(client, owner, cid, result).status_code == 403


def test_ai_company_allowance_includes_failures_and_manual_commit_remains_available(client: TestClient) -> None:
    owner, company, _ = new_company(client)
    cid = company["id"]
    result = preview(client, owner, cid).json()

    async def fill(db: Any) -> None:
        for _ in range(10):
            db.add(
                ProviderRequest(
                    company_id=cid,
                    user_id=uid(owner),
                    provider="mock",
                    model="mock",
                    request_type="business_import_review",
                    ok=False,
                )
            )

    db_call(client, fill)
    assert review(client, owner, cid, result).status_code == 429
    assert commit(client, owner, cid, result).status_code == 200


def test_docx_preview_reports_source_and_refuses_cross_company_writes(client: TestClient) -> None:
    owner, company, _ = new_company(client)
    result = preview(
        client,
        owner,
        company["id"],
        docx_bytes(content()["body"]),
        "policy.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    assert result.status_code == 200, result.text
    other, _, _ = new_company(client)
    assert commit(client, other, company["id"], result.json()).status_code == 404
