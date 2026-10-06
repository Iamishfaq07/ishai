"""India campus mode: the education grade (CGPA or percentage) everywhere it goes, Class 10 and 12 as education, and
the Campus and Harvard templates, which stay single-column and carry their look into Word."""

import io
import json

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from fastapi.testclient import TestClient
from pydantic import ValidationError

from orbit_api.resumes import ai
from orbit_api.resumes.ats import resume_text, score_resume
from orbit_api.resumes.claims import ground_import
from orbit_api.resumes.docx_export import resume_docx
from orbit_api.resumes.schema import Basics, Education, ResumeData, ResumeDesign
from tests.test_api import _register
from tests.test_resumes import strong_resume


def fresher_education() -> list[Education]:
    return [
        Education(school="Vellore Institute of Technology", degree="B.Tech", field="Computer Science", start="2021",
                  end="2025", grade="8.62 / 10 CGPA", details="Merit scholarship"),
        Education(school="Delhi Public School (CBSE)", degree="Class 12 (PCM)", end="2021", grade="94.2%"),
        Education(school="Delhi Public School (CBSE)", degree="Class 10", end="2019", grade="96.4%"),
    ]


# ---- the grade field -----------------------------------------------------------------------------------------


def test_the_grade_is_optional_and_bounded() -> None:
    assert Education().grade == ""
    assert Education(grade="8.6 / 10 CGPA").grade == "8.6 / 10 CGPA"
    with pytest.raises(ValidationError):
        Education(grade="9" * 61)


def test_a_resume_saved_before_grades_existed_still_loads(client: TestClient) -> None:
    old = {"education": [{"id": "e1", "school": "UT Austin", "degree": "BSc", "field": "CS", "start": "2012",
                          "end": "2016", "details": ""}]}
    assert ResumeData.model_validate(old).education[0].grade == ""
    h = _register(client)
    r = client.post("/v1/resumes", headers=h, json={"title": "Old", "data": old})
    assert r.status_code == 201, r.text
    assert client.get(f"/v1/resumes/{r.json()['id']}", headers=h).json()["data"]["education"][0]["grade"] == ""
    # ...and a grade saved now comes back.
    data = ResumeData(education=fresher_education()).model_dump(mode="json")
    r = client.post("/v1/resumes", headers=h, json={"title": "New", "data": data})
    got = client.get(f"/v1/resumes/{r.json()['id']}", headers=h).json()["data"]["education"]
    assert [e["grade"] for e in got] == ["8.62 / 10 CGPA", "94.2%", "96.4%"]


def test_the_grade_is_in_the_resume_text_and_the_ais_view_of_it() -> None:
    data = ResumeData(basics=Basics(name="Aarav"), education=fresher_education())
    assert "8.62 / 10 CGPA" in resume_text(data) and "94.2%" in resume_text(data)
    seen = json.loads(ai._resume_json(data))["education"]
    assert seen[0] == {"degree": "B.Tech, Computer Science", "school": "Vellore Institute of Technology",
                       "grade": "8.62 / 10 CGPA"}
    assert "grade" not in json.loads(ai._resume_json(ResumeData(education=[Education(school="X")])))["education"][0]


# ---- AI import: CGPA and percentages, grounded --------------------------------------------------------------

SOURCE = """Priya Nair
priya.nair@example.com | +91 91234 56789
EDUCATION
National Institute of Technology Calicut - B.Tech, Electronics (2020-2024) - CGPA: 8.71/10
Kendriya Vidyalaya, CBSE - Class XII (Science), 2020 - 91.60%
Kendriya Vidyalaya, CBSE - Class X, 2018 - First Division
"""


def _answer(grades: list[str]) -> str:
    schools = [("National Institute of Technology Calicut", "B.Tech", "Electronics", "2020", "2024"),
               ("Kendriya Vidyalaya, CBSE", "Class XII (Science)", "", "", "2020"),
               ("Kendriya Vidyalaya, CBSE", "Class X", "", "", "2018")]
    return json.dumps({
        "basics": {"name": "Priya Nair", "email": "priya.nair@example.com", "phone": "+91 91234 56789"},
        "education": [{"school": s, "degree": d, "field": f, "start": a, "end": b, "grade": g}
                      for (s, d, f, a, b), g in zip(schools, grades, strict=True)],
    })


def test_the_import_prompt_asks_for_the_grade_and_school_classes() -> None:
    assert '"grade"' in ai.IMPORT_SYSTEM and "CGPA" in ai.IMPORT_SYSTEM and "Class 12" in ai.IMPORT_SYSTEM
    assert "Never work one out" in ai.IMPORT_SYSTEM
    assert "Positions of Responsibility" in ai.IMPORT_SYSTEM and "Internships" in ai.IMPORT_SYSTEM


def test_the_import_reads_grades_and_keeps_the_ones_the_source_gives() -> None:
    data = ai.parse_structured(_answer(["8.71 / 10 CGPA", "91.6%", "First Division"]))
    assert [e.grade for e in data.education] == ["8.71 / 10 CGPA", "91.6%", "First Division"]
    out, notes = ground_import(data, SOURCE)
    # "/ 10" needn't be written in the source; "91.6" is the source's "91.60"; a grade in words is found as words.
    assert [e.grade for e in out.education] == ["8.71 / 10 CGPA", "91.6%", "First Division"]
    assert not any("grade" in n for n in notes)


def test_the_import_blanks_a_grade_the_source_never_gives() -> None:
    data = ai.parse_structured(_answer(["8.9 / 10 CGPA", "92%", "Distinction"]))
    out, notes = ground_import(data, SOURCE)
    assert [e.grade for e in out.education] == ["", "", ""]
    assert "3 grade(s) the source never gives" in notes
    assert out.education[1].degree == "Class XII (Science)"  # the rest of the entry stays


def test_a_long_grade_is_clipped_not_refused() -> None:
    data = ai.parse_structured(_answer(["8.71 " + "x" * 200, "", ""]))
    assert len(data.education[0].grade) == 60 and data.education[1].grade == ""


def test_the_import_endpoint_grounds_grades(client: TestClient, monkeypatch: pytest.MonkeyPatch) -> None:
    from orbit_api.providers.llm.mock import MockLLMProvider
    from orbit_api.resumes import routes

    monkeypatch.setattr(routes, "get_llm_provider",
                        lambda: MockLLMProvider(script=[_answer(["8.71/10 CGPA", "91.6%", "75%"])]))
    h = _register(client)
    r = client.post("/v1/resumes/import", headers=h, json={"source": "text", "text": SOURCE})
    assert r.status_code == 200, r.text
    assert [e["grade"] for e in r.json()["data"]["education"]] == ["8.71/10 CGPA", "91.6%", ""]


# ---- Word export ---------------------------------------------------------------------------------------------


def _paragraphs(data: ResumeData, design: ResumeDesign) -> list[str]:
    return [p.text for p in Document(io.BytesIO(resume_docx(data, design))).paragraphs]


def test_word_puts_the_grade_under_the_degree() -> None:
    data = ResumeData(basics=Basics(name="Aarav Sharma"), education=fresher_education())
    lines = _paragraphs(data, ResumeDesign())
    at = lines.index("B.Tech, Computer Science, Vellore Institute of Technology\t2021 – 2025")
    assert lines[at + 1] == "8.62 / 10 CGPA  ·  Merit scholarship"
    assert "Class 12 (PCM), Delhi Public School (CBSE)\t2021" in lines and "94.2%" in lines
    # No grade or details, no line under the degree.
    plain = _paragraphs(ResumeData(education=[Education(school="UT Austin", degree="BSc", end="2016")]),
                        ResumeDesign())
    assert plain[-2:] == ["EDUCATION", "BSc, UT Austin\t2016"]


def test_word_keeps_the_campus_look() -> None:
    data = strong_resume()
    data.education = fresher_education()
    doc = Document(io.BytesIO(resume_docx(data, ResumeDesign(template="campus", font="cambria"))))
    paras = doc.paragraphs
    lines = [p.text for p in paras]
    # Name and contact line centred, the contacts split by " | ".
    assert paras[0].alignment == WD_ALIGN_PARAGRAPH.CENTER and paras[0].runs[0].font.small_caps
    assert "sam@example.com | +1 555 123 4567 | Austin, TX" in lines[1]
    assert paras[1].alignment == WD_ALIGN_PARAGRAPH.CENTER
    # Small-caps headings in the section's own case.
    heads = [p for p in paras if p.style.name == "Heading 1"]
    assert heads[0].text == "Summary" and all(p.runs[0].font.small_caps for p in heads)
    assert all(p.runs[0].font.name == "Cambria" for p in heads)
    # Company and dates, then the role (and location) in italics; school and dates, then the degree and grade.
    at = lines.index("PayCo\tJan 2020 – Present")
    assert lines[at + 1] == "Senior Engineer"
    assert all(r.italic for r in paras[at + 1].runs)
    at = lines.index("Vellore Institute of Technology\t2021 – 2025")
    assert lines[at + 1] == "B.Tech, Computer Science\t8.62 / 10 CGPA" and lines[at + 2] == "Merit scholarship"
    at = lines.index("Delhi Public School (CBSE)\t2021")
    assert lines[at + 1] == "Class 12 (PCM)\t94.2%"


def test_word_keeps_the_harvard_look() -> None:
    doc = Document(io.BytesIO(resume_docx(strong_resume(), ResumeDesign(template="harvard", font="inter"))))
    heads = [p for p in doc.paragraphs if p.style.name == "Heading 1"]
    assert heads and all(p.alignment == WD_ALIGN_PARAGRAPH.CENTER and p.text.isupper() for p in heads)
    # A serif for the name and headings even next to a sans-serif body.
    assert doc.styles["Normal"].font.name == "Calibri"
    assert heads[0].runs[0].font.name == "Georgia" and doc.paragraphs[0].runs[0].font.name == "Georgia"
    # Word puts a style's theme font ahead of the run's own, so Heading 1 has none left.
    assert doc.styles["Heading 1"].element.rPr.rFonts.get(qn("w:asciiTheme")) is None
    assert " • " in doc.paragraphs[1].text
    # Times stays Times.
    doc = Document(io.BytesIO(resume_docx(strong_resume(), ResumeDesign(template="harvard", font="times"))))
    assert next(p for p in doc.paragraphs if p.style.name == "Heading 1").runs[0].font.name == "Times New Roman"


def test_other_templates_keep_their_word_look() -> None:
    doc = Document(io.BytesIO(resume_docx(strong_resume(), ResumeDesign(template="classic"))))
    assert doc.paragraphs[0].alignment is None
    heads = [p for p in doc.paragraphs if p.style.name == "Heading 1"]
    assert heads[0].text == "SUMMARY" and not heads[0].runs[0].font.small_caps
    assert "Senior Engineer, PayCo\tJan 2020 – Present" in [p.text for p in doc.paragraphs]


# ---- the templates -------------------------------------------------------------------------------------------


@pytest.mark.parametrize("template", ["campus", "harvard"])
def test_campus_and_harvard_are_single_column_and_score_like_classic(template: str) -> None:
    design = ResumeDesign(template=template)  # type: ignore[arg-type]
    report = score_resume(strong_resume(), design)
    assert {c.id: c for c in report.checks}["single_column"].passed
    assert report.score == score_resume(strong_resume(), ResumeDesign()).score


def test_the_api_saves_the_new_templates(client: TestClient) -> None:
    h = _register(client)
    for template in ("campus", "harvard"):
        r = client.post("/v1/resumes", headers=h, json={"title": template, "design": {"template": template}})
        assert r.status_code == 201, r.text
        assert client.get(f"/v1/resumes/{r.json()['id']}", headers=h).json()["design"]["template"] == template
