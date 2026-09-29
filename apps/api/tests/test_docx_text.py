"""Word resumes lay out their content in tables, text boxes and page headers: all of it must be read, once."""

import io

import docx
from docx.oxml import parse_xml

from orbit_api.documents.parsing import extract_text

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
MC = "http://schemas.openxmlformats.org/markup-compatibility/2006"


def _save(d: docx.document.Document) -> bytes:
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_tables_headers_and_footers_are_read_in_order() -> None:
    d = docx.Document()
    d.sections[0].header.paragraphs[0].text = "jane@example.com | +1 555 0100"
    d.add_paragraph("Jane Doe")
    table = d.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Skills"
    table.cell(0, 0).add_paragraph("Kubernetes")
    table.cell(0, 1).text = "Senior Engineer at Acme"
    d.sections[0].footer.paragraphs[0].text = "References on request"
    lines = extract_text(_save(d), "docx").split("\n")
    assert lines == ["jane@example.com | +1 555 0100", "Jane Doe", "Skills", "Kubernetes", "Senior Engineer at Acme",
                     "References on request"]


def test_a_text_box_is_read_once() -> None:
    d = docx.Document()
    par = d.add_paragraph("Before the box")
    # Word stores a text box twice: the modern drawing (mc:Choice) and a legacy copy (mc:Fallback).
    box = (f'<w:r xmlns:w="{W}" xmlns:mc="{MC}"><mc:AlternateContent>'
           '<mc:Choice Requires="wps"><w:txbxContent><w:p><w:r><w:t>Profile: backend engineer</w:t></w:r></w:p>'
           '</w:txbxContent></mc:Choice>'
           '<mc:Fallback><w:txbxContent><w:p><w:r><w:t>Profile: backend engineer</w:t></w:r></w:p>'
           '</w:txbxContent></mc:Fallback></mc:AlternateContent></w:r>')
    par._p.append(parse_xml(box))
    d.add_paragraph("After the box")
    text = extract_text(_save(d), "docx")
    assert text.count("Profile: backend engineer") == 1
    assert text.index("Before the box") < text.index("Profile") < text.index("After the box")
    assert "Before the box\n" in text  # the outer paragraph doesn't swallow the box's text
