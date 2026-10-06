"""The resume as a Word file: one plain column in reading order, built from real Word styles (headings, bullet
lists, a right tab stop for dates), so ATS software and recruiters who ask for .docx read it cleanly.

Only the design choices that survive in Word carry over: font, size, accent colour, spacing and page size, and the
Campus and Harvard templates' headings and two-line entries. The layout is always single-column and has no photo,
since that's what ATS parsers read most reliably.
"""

import io
from datetime import UTC, datetime

from docx import Document
from docx.document import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from docx.text.paragraph import Paragraph

from orbit_api.resumes.schema import EXTRA_PREFIX, ExtraSection, ResumeData, ResumeDesign

SECTION_LABEL: dict[str, str] = {
    "summary": "Summary",
    "experience": "Experience",
    "education": "Education",
    "skills": "Skills",
    "projects": "Projects",
    "certifications": "Certifications",
}
# Fonts every copy of Word has; the web-only ones fall back to the closest match.
WORD_FONT = {"inter": "Calibri", "arial": "Arial", "georgia": "Georgia", "garamond": "Garamond",
             "grotesk": "Calibri", "calibri": "Calibri", "cambria": "Cambria", "helvetica": "Arial",
             "times": "Times New Roman", "trebuchet": "Trebuchet MS"}
SERIF_FONTS = frozenset({"georgia", "garamond", "cambria", "times"})
# Campus and Harvard keep their look in Word: the name and contact line centred, their own headings (Campus in small
# capitals, Harvard centred in a serif), and each entry on two lines: the school or company in bold with the dates
# flush right, the degree or role in italics below with the grade or location flush right.
STACKED = frozenset({"campus", "harvard"})
CONTACT_SEP = {"campus": " | ", "harvard": "  •  "}
LINE_SPACING = {"tight": 1.05, "normal": 1.15, "airy": 1.3}
PAGE_MM = {"a4": (210, 297), "letter": (215.9, 279.4)}
# Page margin, gap before a section heading, gap after an entry: the same steps as the preview.
SPACING = {"compact": (12, 9, 5), "normal": (16, 13, 7), "relaxed": (20, 17, 10)}
# Campus is tighter between sections and entries (the margin stays), as on the preview.
CAMPUS_GAPS = {"compact": (6, 4), "normal": (8, 5), "relaxed": (11, 6)}


def _date_range(start: str, end: str, current: bool = False) -> str:
    to = "Present" if current else end.strip()
    return " – ".join(p for p in (start.strip(), to) if p)


def extra_section(data: ResumeData, key: str) -> ExtraSection | None:
    """The extra section a section_order entry ("extra:<id>") names, if it exists."""
    if not key.startswith(EXTRA_PREFIX):
        return None
    return next((x for x in data.extras if x.id == key[len(EXTRA_PREFIX):]), None)


def _has_content(data: ResumeData, section: str) -> bool:
    if section == "summary":
        return bool(data.basics.summary.strip())
    if section == "skills":
        return any(i.strip() for s in data.skills for i in s.items)
    if section.startswith(EXTRA_PREFIX):
        extra = extra_section(data, section)
        return extra is not None and bool(extra.title.strip()) and any(i.strip() for i in extra.items)
    return bool(getattr(data, section, None))


_PBDR_SUCCESSORS = ("w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
                    "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd",
                    "w:snapToGrid", "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents",
                    "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment", "w:textboxTightWrap",
                    "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange")


def _bottom_border(par: Paragraph, color: str) -> None:
    ppr = par._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", color.lstrip("#"))):
        bottom.set(qn(key), value)
    bdr.append(bottom)
    # Word rejects paragraph properties out of schema order ("unreadable content"), so place it before its successors.
    ppr.insert_element_before(bdr, *_PBDR_SUCCESSORS)


class _Writer:
    def __init__(self, design: ResumeDesign) -> None:
        self.doc: DocxDocument = Document()
        self.accent = RGBColor.from_string(design.accent.lstrip("#").upper())
        self.accent_hex = design.accent
        self.template = design.template
        self.stacked = design.template in STACKED
        # Harvard's name and headings are always a serif: the chosen one, or Georgia next to a sans-serif body.
        serif = design.template != "harvard" or design.font in SERIF_FONTS
        self.title_font = WORD_FONT[design.font] if serif else "Georgia"
        margin, self.section_gap, self.item_gap = SPACING[design.spacing]
        if design.template == "campus":
            self.section_gap, self.item_gap = CAMPUS_GAPS[design.spacing]
        section = self.doc.sections[0]
        width, height = PAGE_MM[design.page]
        section.page_width, section.page_height = Mm(width), Mm(height)
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(section, side, Mm(margin))
        self.text_width = Mm(width - 2 * margin)
        normal = self.doc.styles["Normal"]
        normal.font.name = WORD_FONT[design.font]
        normal.font.size = Pt(design.font_size)
        # East-Asian and complex-script text would otherwise keep the theme font.
        rfonts = normal.element.get_or_add_rPr().get_or_add_rFonts()
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rfonts.set(qn(attr), WORD_FONT[design.font])
        normal.paragraph_format.space_after = Pt(0)
        normal.paragraph_format.line_spacing = LINE_SPACING[design.line_height]
        # Heading 1 comes with theme fonts, which Word puts ahead of the font a run names: without them the headings
        # are set in the resume's font (Harvard's serif), not the theme's.
        heading_fonts = self.doc.styles["Heading 1"].element.get_or_add_rPr().get_or_add_rFonts()
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            heading_fonts.attrib.pop(qn(attr), None)

    def properties(self, kind: str, name: str) -> None:
        """What Word's File > Info shows, and what some job portals read: the candidate, not "python-docx"."""
        props = self.doc.core_properties
        name = name.strip()
        props.title = f"{name} - {kind}" if name else kind
        props.author = props.last_modified_by = name
        props.comments = props.subject = props.keywords = props.category = ""

    def par(self, text: str = "", *, bold: bool = False, size: float | None = None, color: RGBColor | None = None,
            style: str | None = None, after: float = 0) -> Paragraph:
        p = self.doc.add_paragraph(style=style)
        if text:
            run = p.add_run(text)
            run.bold = bold
            if size:
                run.font.size = Pt(size)
            if color is not None:
                run.font.color.rgb = color
        p.paragraph_format.space_after = Pt(after)
        return p

    def heading(self, label: str) -> None:
        # A real Heading 1 so parsers and screen readers see the section structure, styled to match the resume.
        p = self.doc.add_paragraph(style="Heading 1")
        campus = self.template == "campus"
        run = p.add_run(label if campus else label.upper())
        run.font.size = Pt(12 if campus else 11)
        if campus:
            run.font.small_caps = True
        run.font.color.rgb = self.accent
        run.font.name = self.title_font
        run.bold = True
        if self.template == "harvard":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fmt = p.paragraph_format
        fmt.space_before, fmt.space_after, fmt.keep_with_next = Pt(self.section_gap), Pt(3), True
        _bottom_border(p, self.accent_hex)

    def entry_line(self, left: list[tuple[str, bool]], right: str, *, italic: bool = False) -> None:
        """"Title, Company · Location" on the left and the dates flush right, on one line."""
        p = self.doc.add_paragraph()
        p.paragraph_format.tab_stops.add_tab_stop(self.text_width, WD_TAB_ALIGNMENT.RIGHT)
        p.paragraph_format.keep_with_next = True
        for text, bold in left:
            run = p.add_run(text)
            run.bold, run.italic = bold, italic
        if right:
            p.add_run("\t" + right).italic = italic

    def stacked_entry(self, top: str, dates: str, sub: str, side: str) -> None:
        """Campus and Harvard: "Company<tab>dates" with the company in bold, then "Role<tab>location" in italics."""
        self.entry_line([(top, True)], dates)
        if sub or side:
            self.entry_line([(sub, False)], side, italic=True)

    def bullets(self, items: list[str]) -> None:
        for item in items:
            if item.strip():
                self.par(item.strip(), style="List Bullet")

    def gap(self) -> None:
        paragraphs = self.doc.paragraphs
        if paragraphs:
            paragraphs[-1].paragraph_format.space_after = Pt(self.item_gap)


def resume_docx(data: ResumeData, design: ResumeDesign) -> bytes:
    w = _Writer(design)
    data = data.visible()
    b = data.basics
    header = [w.par(b.name.strip() or "Your name", bold=True, size=design.font_size * 2)]
    header[0].runs[0].font.name = w.title_font
    if w.template == "campus":
        header[0].runs[0].font.small_caps = True
    if b.headline.strip():
        header.append(w.par(b.headline.strip(), color=w.accent, size=design.font_size * 1.1))
    contacts = [s.strip() for s in (b.email, b.phone, b.location, *(link.url for link in b.links)) if s.strip()]
    if contacts:
        header.append(w.par(CONTACT_SEP.get(w.template, "  ·  ").join(contacts)))
    if w.stacked:
        for p in header:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for section in data.section_order:
        if not _has_content(data, section):
            continue
        extra = extra_section(data, section)
        w.heading(extra.title.strip() if extra else data.titles.get(section) or SECTION_LABEL[section])
        if extra:
            items = [i.strip() for i in extra.items if i.strip()]
            if extra.style == "inline":
                w.par("  ·  ".join(items))
            else:
                w.bullets(items)
        elif section == "summary":
            w.par(b.summary.strip())
        elif section == "experience":
            for e in data.experience:
                if w.stacked:
                    w.stacked_entry(e.company or e.title, _date_range(e.start, e.end, e.current),
                                    e.title if e.company else "", e.location)
                    w.bullets(e.bullets)
                    w.gap()
                    continue
                left = [(e.title, True)] if e.title else []
                if e.company:
                    left.append(((", " if e.title else "") + e.company, False))
                if e.location:
                    left.append((" · " + e.location, False))
                w.entry_line(left, _date_range(e.start, e.end, e.current))
                w.bullets(e.bullets)
                w.gap()
        elif section == "education":
            for ed in data.education:
                degree = ", ".join(p for p in (ed.degree, ed.field) if p)
                if w.stacked:
                    w.stacked_entry(ed.school or degree, _date_range(ed.start, ed.end), degree if ed.school else "",
                                    ed.grade.strip())
                    if ed.details.strip():
                        w.par(ed.details.strip())
                    w.gap()
                    continue
                left = [(degree or ed.school, True)]
                if degree and ed.school:
                    left.append((", " + ed.school, False))
                w.entry_line(left, _date_range(ed.start, ed.end))
                # The grade leads the line under the degree ("8.6 / 10 CGPA · Dean's list"), as on the preview.
                below = "  ·  ".join(x for x in (ed.grade.strip(), ed.details.strip()) if x)
                if below:
                    w.par(below)
                w.gap()
        elif section == "skills":
            for group in data.skills:
                items = [i.strip() for i in group.items if i.strip()]
                if items:
                    p = w.par()
                    if group.name.strip():
                        p.add_run(group.name.strip() + ": ").bold = True
                    p.add_run(", ".join(items))
        elif section == "projects":
            for pr in data.projects:
                p = w.par(pr.name, bold=True)
                p.paragraph_format.keep_with_next = True
                if pr.link.strip():
                    p.add_run(" · " + pr.link.strip())
                if pr.description.strip():
                    w.par(pr.description.strip())
                w.bullets(pr.bullets)
                w.gap()
        elif section == "certifications":
            for c in data.certifications:
                p = w.par(c.name, bold=True)
                issued = ", ".join(x for x in (c.issuer, c.date) if x.strip())
                if issued:
                    p.add_run(" · " + issued)

    w.properties("Resume", b.name)
    buf = io.BytesIO()
    w.doc.save(buf)
    return buf.getvalue()


def letter_docx(text: str, data: ResumeData, design: ResumeDesign) -> bytes:
    """The cover letter on a page that matches the resume: the same name and contact line, then today's date and the
    letter's paragraphs (a single line break inside one, e.g. "Sincerely," / the name, stays a line break)."""
    w = _Writer(design)
    b = data.basics
    if b.name.strip():
        w.par(b.name.strip(), bold=True, size=design.font_size * 1.6)
    contacts = [s.strip() for s in (b.email, b.phone, b.location, *(link.url for link in b.links)) if s.strip()]
    if contacts:
        w.par("  ·  ".join(contacts), after=w.section_gap)
    today = datetime.now(UTC)
    w.par(f"{today.day} {today:%B %Y}", after=w.section_gap)
    for block in (p.strip() for p in text.replace("\r\n", "\n").split("\n\n")):
        if not block:
            continue
        p = w.par(after=w.item_gap + 3)
        for i, line in enumerate(block.split("\n")):
            if i:
                p.add_run().add_break()
            p.add_run(line.strip())
    w.properties("Cover letter", b.name)
    buf = io.BytesIO()
    w.doc.save(buf)
    return buf.getvalue()
