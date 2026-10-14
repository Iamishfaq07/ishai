"""Safe document text extraction + chunking into retrievable facts (spec §20–§22).

All uploaded text is untrusted: we only extract plain text, never execute or follow anything in it,
and the prompt layer fences it as reference data.

Hostile files are handled in two layers:
- A DOCX is a ZIP: its central directory is inspected *before* any parser sees it (entry count, per-entry and
  total uncompressed size, compression ratio, path traversal), so a ZIP bomb is refused without inflating it.
- PDF/DOCX extraction runs in a short-lived child process (`extract_text_isolated`) with a timeout and, on Linux,
  memory and CPU limits. A parser that hangs, balloons or crashes takes down the child, never the API.

A PDF is read by pypdf, then by PDFium if pypdf finds no text (fonts pypdf can't map, some design-tool exports),
then, for a scanned PDF with no text layer, by OCR: PDFium renders the first pages and Tesseract reads them.
"""

import io
import json
import logging
import multiprocessing
import os
import re
import shutil
import subprocess
import sys
import threading
import zipfile
from dataclasses import dataclass
from multiprocessing.connection import Connection
from pathlib import PurePosixPath
from typing import Any

from orbit_api.documents.uploads import sanitize_text

MAX_PAGES = 40
MAX_CHARS = 200_000

# DOCX archive limits (a real resume is well under 1 MB and a few dozen entries).
MAX_ZIP_ENTRIES = 1000
MAX_ZIP_ENTRY_BYTES = 25 * 1024 * 1024
MAX_ZIP_TOTAL_BYTES = 50 * 1024 * 1024
MAX_ZIP_RATIO = 100  # uncompressed / compressed, checked on entries big enough to matter
ZIP_RATIO_MIN_BYTES = 1024 * 1024
# Child-process limits for PDF/DOCX extraction. OCR of a scanned PDF is slow on a small CPU, so PDFs get longer.
PARSE_TIMEOUT_S = 20.0
PDF_PARSE_TIMEOUT_S = 75.0
# OCR of scanned PDFs: pages read, render resolution (Tesseract is most accurate around 300 dpi), per-page cap.
OCR_MAX_PAGES = 3
OCR_DPI = 300
OCR_PAGE_TIMEOUT_S = 20.0
OCR_MAX_SIDE_PX = 4000
PARSE_MEMORY_BYTES = 768 * 1024 * 1024
PARSE_CPU_SECONDS = 30
# Parser children at once, server-wide: each may take hundreds of MB, so a burst of uploads queues instead of running
# the instance out of memory. A wait past PARSE_QUEUE_S answers "busy" rather than holding the request.
PARSE_CONCURRENCY = 2
PARSE_QUEUE_S = 30.0
_parsers = threading.BoundedSemaphore(PARSE_CONCURRENCY)
MAX_CHILD_REPLY_BYTES = 8 * 1024 * 1024  # the extracted text, as JSON
# Set by `isolation_selftest()` at boot if child-process parsing doesn't work on this host; extraction then runs
# in-process (archive inspection still applies) instead of failing every upload.
ISOLATION_BROKEN = False
log = logging.getLogger(__name__)

ALLOWED_TYPES = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "text/plain": "txt",
    "text/markdown": "md",
}
ALLOWED_EXT = {".pdf": "pdf", ".docx": "docx", ".txt": "txt", ".md": "md"}
ALLOWED_EXT.update(dict.fromkeys((".py", ".js", ".ts", ".tsx", ".jsx", ".rs", ".go", ".java", ".c", ".cpp",
                                 ".h", ".sql", ".sh", ".yaml", ".yml", ".json"), "code"))


class ParseError(ValueError):
    pass


def detect_format(filename: str, content_type: str | None, head: bytes) -> str:
    ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    fmt = ALLOWED_EXT.get(ext) or ALLOWED_TYPES.get((content_type or "").split(";")[0].strip())
    if fmt is None:
        raise ParseError("Unsupported file type. Upload PDF, DOCX, TXT or MD.")
    mime = (content_type or "").split(";", 1)[0].strip()
    document_mimes = {"pdf": {"application/pdf", "application/octet-stream"},
                      "docx": {"application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                               "application/zip", "application/octet-stream"}}
    if fmt in document_mimes and mime and mime not in document_mimes[fmt]:
        raise ParseError("File content does not match its MIME type or extension.")
    if (mime in ALLOWED_TYPES and ALLOWED_TYPES[mime] != fmt
            and not (fmt in {"txt", "md", "code"} and mime.startswith("text/"))):
        raise ParseError("File content does not match its MIME type or extension.")
    if fmt == "code" and mime and not (mime.startswith("text/") or mime in {
        "application/json", "application/javascript", "application/octet-stream", "application/x-python-code"}):
        raise ParseError("Code uploads must be plain text.")
    # Magic-byte check so a renamed binary can't masquerade as another format.
    if fmt == "pdf" and not head.startswith(b"%PDF"):
        raise ParseError("File is not a valid PDF.")
    if fmt == "docx" and not head.startswith(b"PK"):
        raise ParseError("File is not a valid DOCX.")
    if fmt in ("txt", "md", "code") and b"\x00" in head:
        raise ParseError("Text file contains binary data.")
    return fmt


def inspect_docx_archive(data: bytes) -> None:
    """Refuse ZIP bombs and odd archives from the central directory alone (nothing is decompressed here)."""
    try:
        zf = zipfile.ZipFile(io.BytesIO(data))
    except (zipfile.BadZipFile, ValueError) as e:
        raise ParseError("File is not a valid DOCX.") from e
    with zf:
        infos = zf.infolist()
        if len(infos) > MAX_ZIP_ENTRIES:
            raise ParseError("This DOCX has too many parts to process safely.")
        total = 0
        names = set()
        for info in infos:
            name = info.filename
            path = PurePosixPath(name.replace("\\", "/"))
            if name.startswith(("/", "\\")) or ".." in path.parts or ":" in name:
                raise ParseError("This DOCX contains an invalid file path.")
            if info.file_size > MAX_ZIP_ENTRY_BYTES:
                raise ParseError("This DOCX is too large to process (an embedded part is over 25 MB).")
            if info.file_size >= ZIP_RATIO_MIN_BYTES and info.file_size > MAX_ZIP_RATIO * max(info.compress_size, 1):
                raise ParseError("This DOCX looks like a compression bomb and was refused.")
            total += info.file_size
            names.add(name)
        if total > MAX_ZIP_TOTAL_BYTES:
            raise ParseError("This DOCX is too large to process (over 50 MB uncompressed).")
        if "word/document.xml" not in names or "[Content_Types].xml" not in names:
            raise ParseError("File is not a valid DOCX.")


def _limit_child_resources() -> None:
    if sys.platform.startswith("linux"):
        import resource

        resource.setrlimit(resource.RLIMIT_AS, (PARSE_MEMORY_BYTES, PARSE_MEMORY_BYTES))
        resource.setrlimit(resource.RLIMIT_CPU, (PARSE_CPU_SECONDS, PARSE_CPU_SECONDS))


def _reply(conn: Connection, status: str, value: str) -> None:
    # JSON bytes, not pickle: a parser compromised by a hostile file can't hand the server an object to unpickle.
    conn.send_bytes(json.dumps([status, value]).encode())


def _extract_in_child(data: bytes, fmt: str, conn: Connection) -> None:
    try:
        _limit_child_resources()
        _reply(conn, "ok", extract_text(data, fmt))
    except ParseError as e:
        _reply(conn, "error", str(e))
    except MemoryError:
        _reply(conn, "error", "This document is too large to process.")
    except BaseException:
        _reply(conn, "error", "Could not read this document.")
    finally:
        conn.close()


def extract_text_isolated(data: bytes, fmt: str) -> str:
    """`extract_text` in a child process with a timeout (plain text needs no parser, so it stays in-process).
    Blocking: call it from a worker thread."""
    if fmt == "docx":
        inspect_docx_archive(data)
    if fmt not in ("pdf", "docx") or ISOLATION_BROKEN:
        return extract_text(data, fmt)
    if not _parsers.acquire(timeout=PARSE_QUEUE_S):
        raise ParseError("Orbit is busy reading other documents. Try again in a minute.")
    try:
        return _parse_in_child(data, fmt)
    finally:
        _parsers.release()


def _parse_in_child(data: bytes, fmt: str) -> str:
    ctx = multiprocessing.get_context("spawn")  # never fork a threaded server
    receiver, sender = ctx.Pipe(duplex=False)
    proc = ctx.Process(target=_extract_in_child, args=(data, fmt, sender), daemon=True)
    proc.start()
    sender.close()
    try:
        if not receiver.poll(PDF_PARSE_TIMEOUT_S if fmt == "pdf" else PARSE_TIMEOUT_S):
            raise ParseError("This document took too long to process.")
        status, value = json.loads(receiver.recv_bytes(MAX_CHILD_REPLY_BYTES))
    except (EOFError, OSError, json.JSONDecodeError, TypeError) as e:  # the child died (limits, crash), or said junk
        raise ParseError("This document couldn't be processed — it may be malformed or too complex.") from e
    finally:
        proc.join(timeout=1)
        if proc.is_alive():
            proc.kill()
            proc.join(timeout=5)
        receiver.close()
    if status != "ok":
        raise ParseError(value)
    return str(value)


def isolation_selftest() -> bool:
    """Parse a tiny generated DOCX through the child process. Run once at boot (in a thread); on failure, log
    loudly and fall back to in-process parsing rather than rejecting every PDF/DOCX upload."""
    global ISOLATION_BROKEN
    try:
        import docx

        d = docx.Document()
        d.add_paragraph("orbit parser self-test")
        buf = io.BytesIO()
        d.save(buf)
        ok = "self-test" in extract_text_isolated(buf.getvalue(), "docx")
    except Exception:
        log.exception("isolated document parsing self-test crashed")
        ok = False
    if not ok:
        ISOLATION_BROKEN = True
        log.critical("isolated document parsing is unavailable on this host; parsing in-process "
                     "(archive limits still apply)")
    return ok


def _pypdf_text(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        raise ParseError("Encrypted PDFs are not supported.")
    return "\n".join((p.extract_text() or "") for p in reader.pages[:MAX_PAGES])


def _pdfium_text(data: bytes) -> str:
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(data)
    try:
        pages = []
        for i in range(min(len(pdf), MAX_PAGES)):
            page = pdf[i]
            textpage = page.get_textpage()
            pages.append(textpage.get_text_bounded())
            textpage.close()
            page.close()
        return "\n".join(pages)
    finally:
        pdf.close()


def ocr_available() -> bool:
    return shutil.which("tesseract") is not None


def _ocr_pdf(data: bytes) -> str:
    """Read a scanned PDF: render its first pages with PDFium and run Tesseract on each (as a subprocess, so it
    inherits this child's memory and CPU limits)."""
    import pypdfium2 as pdfium

    exe = shutil.which("tesseract")
    if exe is None:
        raise ParseError("This PDF is a scanned image with no text in it, and this server can't read images. "
                         "Export it again as a text PDF, or paste the text instead.")
    env = {**os.environ, "OMP_THREAD_LIMIT": "1"}  # one thread: steadier on small shared CPUs
    pdf = pdfium.PdfDocument(data)
    try:
        texts = []
        for i in range(min(len(pdf), OCR_MAX_PAGES)):
            page = pdf[i]
            buf = io.BytesIO()
            scale = min(OCR_DPI / 72, OCR_MAX_SIDE_PX / max(*page.get_size(), 1))  # a poster-sized page stays small
            page.render(scale=scale, grayscale=True).to_pil().save(buf, format="PNG")
            page.close()
            try:
                # Fixed arguments; the untrusted file only reaches Tesseract as rendered PNG bytes on stdin.
                out = subprocess.run([exe, "stdin", "stdout", "-l", "eng", "--psm", "3"], input=buf.getvalue(),  # noqa: S603
                                     capture_output=True, timeout=OCR_PAGE_TIMEOUT_S, env=env, check=False)
            except subprocess.TimeoutExpired as e:
                raise ParseError("This scanned PDF took too long to read. Paste the text instead.") from e
            texts.append(out.stdout.decode("utf-8", errors="replace"))
        return "\n".join(texts)
    finally:
        pdf.close()


def _pdf_text(data: bytes) -> str:
    try:
        text = _pypdf_text(data)
    except ParseError:
        raise
    except Exception:  # pypdf chokes on some files PDFium reads fine
        log.info("pypdf couldn't read a PDF; trying PDFium")
        text = ""
    if not text.strip():
        text = _pdfium_text(data)
    if not text.strip():
        text = _ocr_pdf(data)
    return text


_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_MC_FALLBACK = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Fallback"


def _docx_lines(root: Any) -> list[str]:
    """Every paragraph under `root` in reading order, including those in tables and text boxes. A text box sits
    inside a run of an outer paragraph, so each piece of text counts only toward its nearest paragraph; the legacy
    copy Word keeps of each text box (mc:Fallback) is skipped so nothing is read twice."""
    lines = []
    for par in root.iter(_W + "p"):
        if any(a.tag == _MC_FALLBACK for a in par.iterancestors()):
            continue
        parts = []
        for node in par.iter(_W + "t", _W + "tab", _W + "br", _W + "cr"):
            # Only text inside this paragraph's own runs: a text box has its own paragraphs, and a paragraph's
            # settings list its tab stops as w:tab too.
            if next(node.iterancestors(_W + "p"), None) is not par or node.getparent().tag != _W + "r":
                continue
            parts.append(node.text or "" if node.tag == _W + "t" else "\t" if node.tag == _W + "tab" else "\n")
        lines.append("".join(parts))
    return lines


def _docx_text(data: bytes) -> str:
    """Body text plus page headers and footers (resume templates often put the contact line there, and lay the
    body out in tables), headers first and footers last."""
    import docx

    d = docx.Document(io.BytesIO(data))
    head: list[str] = []
    foot: list[str] = []
    for part in d.part.package.iter_parts():
        name = str(part.partname)
        target = head if name.startswith("/word/header") else foot if name.startswith("/word/footer") else None
        if target is not None and hasattr(part, "element"):
            target.extend(line for line in _docx_lines(part.element) if line.strip() and line not in target)
    return "\n".join([*head, *_docx_lines(d.element.body), *foot])


def extract_text(data: bytes, fmt: str) -> str:
    try:
        if fmt == "pdf":
            text = _pdf_text(data)
        elif fmt == "docx":
            text = _docx_text(data)
        else:
            if b"\x00" in data:
                raise ParseError("Text file contains binary data.")
            text = data.decode("utf-8")
    except ParseError:
        raise
    except Exception as e:
        raise ParseError("Could not read this document.") from e
    text = sanitize_text(text)
    if not text.strip():
        raise ParseError("No readable text found in this file. Paste the text instead.")
    return text[:MAX_CHARS]


@dataclass
class ChunkDraft:
    category: str
    label: str
    text: str
    verified: bool


_RESUME_SECTIONS: list[tuple[str, tuple[str, ...]]] = [
    ("summary", ("summary", "profile", "about", "objective")),
    ("experience", ("experience", "employment", "work history", "professional experience", "career")),
    ("project", ("projects", "selected projects", "portfolio")),
    ("skill", ("skills", "technical skills", "technologies", "tools", "core competencies")),
    ("education", ("education", "academic")),
    ("certification", ("certifications", "certificates", "licenses")),
    ("achievement", ("achievements", "awards", "honors", "publications")),
]
_JD_SECTIONS: list[tuple[str, tuple[str, ...]]] = [
    ("company", ("about us", "about the company", "who we are", "our mission", "about the team")),
    ("responsibility", ("responsibilities", "what you'll do", "what you will do", "the role", "your role")),
    ("requirement", ("requirements", "qualifications", "what you bring", "must have", "you have",
                     "what we're looking for", "preferred", "nice to have", "skills")),
]
_BULLET = re.compile(r"^\s*(?:[-*•▪◦·]|\d+[.)])\s*")


def _section_of(line: str, table: list[tuple[str, tuple[str, ...]]]) -> str | None:
    norm = re.sub(r"[^a-z' ]", "", line.lower()).strip()
    if not norm or len(norm) > 40:
        return None
    for cat, names in table:
        if norm in names:
            return cat
    return None


def _structured_chunks(
    text: str, table: list[tuple[str, tuple[str, ...]]], default: str, verified: bool, label_prefix: str
) -> list[ChunkDraft]:
    chunks: list[ChunkDraft] = []
    section = default
    heading = ""  # current employer/role line inside experience
    for raw in text.split("\n"):
        line = raw.strip()
        if not line:
            continue
        sec = _section_of(line, table)
        if sec:
            section, heading = sec, ""
            continue
        is_bullet = bool(_BULLET.match(raw))
        body = _BULLET.sub("", line).strip()
        # Skills can be one or two characters (Go, R, C, C#, JS): keep them there; elsewhere drop fragments.
        if len(body) < (1 if section == "skill" else 3) or not any(ch.isalnum() for ch in body):
            continue
        if not is_bullet and len(body) <= 90 and section in ("experience", "project"):
            heading = body  # e.g. "Security Engineer — Acme Corp (2021–2024)"
            continue
        full = f"{heading}: {body}" if heading else body
        label = f"{label_prefix}: {(heading or section.title())[:80]}"
        chunks.append(ChunkDraft(section, label, full[:1200], verified))
    return chunks


def _excerpt_chunks(text: str, label_prefix: str, size: int = 700) -> list[ChunkDraft]:
    paras = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    chunks: list[ChunkDraft] = []
    buf = ""
    for p in paras:
        if len(buf) + len(p) > size and buf:
            chunks.append(ChunkDraft("excerpt", f"{label_prefix}: {buf[:50]}…", buf, False))
            buf = ""
        buf = f"{buf}\n\n{p}".strip()
    if buf:
        chunks.append(ChunkDraft("excerpt", f"{label_prefix}: {buf[:50]}…", buf, False))
    return chunks


def chunk_document(text: str, kind: str, title: str) -> list[ChunkDraft]:
    if kind == "resume":
        return _structured_chunks(text, _RESUME_SECTIONS, "summary", verified=True, label_prefix="resume")
    if kind == "job_description":
        return _structured_chunks(text, _JD_SECTIONS, "requirement", verified=False, label_prefix="JD")
    return _excerpt_chunks(text, label_prefix=title[:40])
